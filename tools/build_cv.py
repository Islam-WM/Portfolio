from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "Islam_Mansour_CV.docx"

NAVY = RGBColor(7, 32, 51)
TEAL = RGBColor(0, 122, 116)
GRAY = RGBColor(83, 96, 105)
LIGHT_GRAY = "C9D2D7"
BLACK = RGBColor(25, 30, 34)
FONT = "Arial"


def set_run_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_bottom_border(paragraph, color=LIGHT_GRAY, size="6", space="3"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_section_heading(doc, text):
    p = doc.add_paragraph(style="Section Heading")
    p.add_run(text.upper())
    set_bottom_border(p, color="7BBAB5", size="5", space="3")
    keep_with_next(p)
    return p


def add_entry(doc, title, stack, description, date=None):
    p = doc.add_paragraph(style="Entry Title")
    if date:
        p.paragraph_format.tab_stops.add_tab_stop(Cm(18.3))
    title_run = p.add_run(title)
    set_run_font(title_run, size=9.7, color=NAVY, bold=True)
    stack_run = p.add_run(f"  |  {stack}")
    set_run_font(stack_run, size=8.4, color=GRAY, italic=True)
    if date:
        date_run = p.add_run(f"\t{date}")
        set_run_font(date_run, size=8.5, color=GRAY, italic=True)
    keep_with_next(p)

    body = doc.add_paragraph(style="Entry Body")
    body.add_run(description)
    return body


def add_labeled_line(doc, label, value):
    p = doc.add_paragraph(style="Compact Line")
    label_run = p.add_run(f"{label}: ")
    set_run_font(label_run, size=8.8, color=NAVY, bold=True)
    value_run = p.add_run(value)
    set_run_font(value_run, size=8.8, color=BLACK)
    return p


def create_document():
    doc = Document()
    section = doc.sections[0]
    # Named override for European recruiting: A4 and compact margins.
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.15)
    section.left_margin = Cm(1.45)
    section.right_margin = Cm(1.45)
    section.header_distance = Cm(0.5)
    section.footer_distance = Cm(0.5)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(9)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(2.6)
    normal.paragraph_format.line_spacing = 1.04

    def make_style(name, size, color, bold=False, italic=False, before=0, after=0, line=1.0):
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = bold
        style.font.italic = italic
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = line
        return style

    make_style("Section Heading", 10.2, TEAL, bold=True, before=7.5, after=4, line=1.0)
    make_style("Entry Title", 9.7, NAVY, bold=True, before=2.5, after=1, line=1.0)
    make_style("Entry Body", 8.8, BLACK, before=0, after=3.4, line=1.05)
    make_style("Compact Line", 8.8, BLACK, before=0, after=1.8, line=1.0)

    doc.core_properties.title = "Islam Mansour - Applied Data Science CV"
    doc.core_properties.subject = "CV for data science, machine learning and data engineering roles"
    doc.core_properties.author = "Islam Mansour"
    doc.core_properties.keywords = "Data Science, Machine Learning, Python, SQL, FastAPI, ETL"

    name = doc.add_paragraph()
    name.paragraph_format.space_before = Pt(0)
    name.paragraph_format.space_after = Pt(0)
    r = name.add_run("ISLAM MANSOUR")
    set_run_font(r, size=25, color=NAVY, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(3)
    r = title.add_run("APPLIED DATA SCIENCE STUDENT  |  MACHINE LEARNING & DATA ENGINEERING")
    set_run_font(r, size=9.2, color=TEAL, bold=True)

    contact = doc.add_paragraph()
    contact.paragraph_format.space_before = Pt(0)
    contact.paragraph_format.space_after = Pt(5)
    contact.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = contact.add_run("Hamburg, Germany  |  +49 176 34249274  |  islamwm05@outlook.de")
    set_run_font(r, size=8.7, color=GRAY)
    set_bottom_border(contact, color="07334D", size="9", space="5")

    summary = doc.add_paragraph()
    summary.paragraph_format.space_before = Pt(3)
    summary.paragraph_format.space_after = Pt(2)
    summary.paragraph_format.line_spacing = 1.08
    r = summary.add_run(
        "Applied Data Science student ranked 6th in the degree program at TU Hamburg. "
        "Builds end-to-end products across machine learning, ETL, backend APIs, media processing and interactive analytics, "
        "with a focus on reliable implementation, validation and clear business-facing outputs."
    )
    set_run_font(r, size=9.2, color=BLACK)

    add_section_heading(doc, "Education")
    add_entry(
        doc,
        "B.Sc. Applied Data Science — Hamburg University of Technology (TU Hamburg)",
        "Ranked 6th in the degree program",
        "Focus: machine learning, statistics, databases and data engineering. Relevant coursework includes Big Data, algorithms and data structures, and statistics.",
        "2022–2026",
    )

    add_section_heading(doc, "Selected Projects")
    add_entry(
        doc,
        "Production-Style Churn Prediction API",
        "Python · Scikit-learn · FastAPI · Pydantic",
        "Deployed a trained gradient-boosting model as a documented REST API. Added schema validation, health-check endpoints, batch prediction and model serialization.",
    )
    add_entry(
        doc,
        "ETL Pipeline & Data Warehouse",
        "Python · SQL · Pandas · SQLite",
        "Integrated CSV, JSON and SQLite sources into staging and mart layers. Implemented CTEs, window functions, automated SQL views and six passing data-quality tests covering nulls, uniqueness and referential integrity.",
    )
    add_entry(
        doc,
        "Interactive Sales Dashboard",
        "Streamlit · Plotly · SQLite · Flask",
        "Built a dashboard with cached SQL access, KPI cards and live filters for region, month and category.",
    )
    add_entry(
        doc,
        "Chess API with AI",
        "FastAPI · Pydantic · SQLite · Minimax · Stockfish",
        "Implemented legal-move validation, persistent PvP games, move history and selectable minimax, random or Stockfish opponents.",
    )
    add_entry(
        doc,
        "Git-Style Repository API",
        "REST API · Push/Pull Workflows",
        "Created API-driven push and pull workflows for transferring and synchronizing project content.",
    )
    add_entry(
        doc,
        "Media Optimization API",
        "REST API · Image & Video Processing",
        "Built automated upload and transformation endpoints that optimize images and selected video formats.",
    )
    add_entry(
        doc,
        "Virtual GPU Renderer",
        "Python · NumPy · Rasterization · Shaders",
        "Implemented a software GPU pipeline with programmable shaders, depth testing, textures and Phong lighting.",
    )
    add_entry(
        doc,
        "Data Cleaning & EDA",
        "Python · Pandas · Seaborn · Matplotlib",
        "Resolved missing values, duplicates, category inconsistencies and outliers in real-estate data, then produced statistical and visual analysis.",
    )

    foundation = doc.add_paragraph(style="Entry Body")
    lead = foundation.add_run("Additional projects: ")
    set_run_font(lead, size=8.8, color=NAVY, bold=True)
    rest = foundation.add_run("NLP sentiment analysis, four-model churn benchmark, automated PDF reporting, C++ ray tracing, animated 3D chess and C terminal games.")
    set_run_font(rest, size=8.8, color=BLACK)

    add_section_heading(doc, "Technical Skills")
    add_labeled_line(doc, "Programming", "Python, SQL, C++, C")
    add_labeled_line(doc, "Data & ML", "Pandas, NumPy, Scikit-learn, Gradient Boosting, Random Forest, NLP, TF-IDF, Matplotlib, Seaborn, Plotly")
    add_labeled_line(doc, "Engineering & Tools", "REST APIs, ETL pipelines, Pydantic, SQLite, FastAPI, Flask, Streamlit, Git, media processing, software rendering, Linux")

    add_section_heading(doc, "Languages")
    add_labeled_line(doc, "Arabic", "Native     |     German: Fluent     |     English: Fluent")

    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.widow_control = True

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    create_document()
